# -*- coding: utf-8 -*-
"""Синхронизация журнала замен с Excel-файлом.

Каждому рабочему ГРП соответствует лист «Ремонт <тип ГРП>»
в файле «Замены.xlsx» (рядом с приложением), по образцу файла
«Лида.xlsx» (листы «Ремонт Лида», «Ремонт Вороново», «Ремонт Кореличи»).

Замены хранятся в PostgreSQL, Excel-файл пересоздаётся из БД при каждом
изменении, поэтому посторонние листы файла сохраняются.
"""
import os

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

# Колонки журнала замен — как в образце (листы «Ремонт …» в Лида.xlsx)
REPLACEMENT_HEADERS = [
    "Дата замены запасной части",
    "Номер запчасти",
    "Тип оборудования (Регулятор, ПСК, ПЗК, Фильтр)",
    "Модель оборудования",
    "Производитель оборудования",
    "Тип объекта (ГРП, ШРП)",
    "Номер объекта",
    "Вид работ при замене запасной части",
    "Причина замены (замена по графику или описание неисправности)",
    "ФИО руководителя газоопасных работ",
]

COLUMN_WIDTHS = [16, 20, 14, 16, 20, 12, 10, 26, 40, 24]

HEADER_FILL = PatternFill("solid", fgColor="DDEBF7")


def default_filename() -> str:
    """Путь к файлу «Замены.xlsx» рядом с модулем приложения."""
    return os.path.join(os.path.dirname(os.path.abspath(__file__)), "Замены.xlsx")


def sheet_name_for(grp_type: str) -> str:
    """Имя листа Excel для ГРП (санитарная очистка + лимит 31 симв.)."""
    invalid = set('[]:*?/\\')
    clean = "".join(ch if ch not in invalid else '_' for ch in str(grp_type)).strip() or "ГРП"
    return f"Ремонт {clean}"[:31]


def _is_catalog_grp(grp_type: str) -> bool:
    return bool(grp_type and 'каталог' in grp_type.lower())


def sync_replacements_workbook(db, filename: str = None) -> str:
    """Пересоздать листы «Ремонт …» по данным БД и сохранить файл."""
    filename = filename or default_filename()

    if os.path.exists(filename):
        wb = load_workbook(filename)
    else:
        wb = Workbook()
        wb.remove(wb.active)

    # Удаляем наши листы «Ремонт ...» и пересоздаём их заново из БД
    for ws in list(wb.worksheets):
        if ws.title.startswith("Ремонт "):
            wb.remove(ws)

    used_titles = set()
    for grp in db.get_all_grp():
        grp_id, grp_type = grp[0], grp[1]
        if _is_catalog_grp(grp_type):
            continue

        title = sheet_name_for(grp_type)
        if title in used_titles:
            title = f"{title[:27]} ({grp_id})"[:31]
        used_titles.add(title)

        ws = wb.create_sheet(title=title)
        ws.append(REPLACEMENT_HEADERS)

        for col in range(1, len(REPLACEMENT_HEADERS) + 1):
            cell = ws.cell(row=1, column=col)
            cell.font = Font(bold=True)
            cell.fill = HEADER_FILL
            cell.alignment = Alignment(wrap_text=True, vertical="top")

        for r in db.get_replacements_by_grp(grp_id):
            # (id, replace_date, part_number, equipment_type, model,
            #  manufacturer, work_type, reason, supervisor)
            ws.append([
                r[1], r[2], r[3], r[4], r[5], grp_type, grp_id, r[6], r[7], r[8],
            ])

        for i, width in enumerate(COLUMN_WIDTHS, start=1):
            ws.column_dimensions[get_column_letter(i)].width = width
        ws.freeze_panes = "A2"

    wb.save(filename)
    return filename