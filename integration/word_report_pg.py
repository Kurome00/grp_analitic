# -*- coding: utf-8 -*-
"""Word-отчёт по ГРП: сведения о ГРП и форма 6.2 расчёта остаточного ресурса.

Отчёт — приложение к паспорту ГРП, поэтому в нём только то, что переносится в
паспорт:

* титульный лист;
* раздел 1 — «Сведения о ГРП»: таблица по полям, которые пользователь ввёл в
  карточке ГРП (см. core.grp_passport);
* раздел 2 — форма 6.2 «Результаты расчёта остаточного ресурса» по двум
  методикам, «Алгоритм 3» и «Алгоритм 4» окна расчёта;
* подписи исполнителей.

Форму 6.2 собирает чистая функция form62_rows: она же проверяется тестами,
поэтому числа в отчёте не могут разойтись с тем, что пользователь видит в
окне «Алгоритмы».
"""
from datetime import date, datetime
from typing import Dict, List, Optional

from docx import Document
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

from core.grp_passport import report_rows
from core.lifetimes import expiry_date, parse_date
from core.timefmt import month_year_text
from logic.algorithms import (
    ALGORITHM_TITLES,
    DAMAGE_MARKERS,
    FAILURE_MARKERS,
    POOR_REPAIR_MARKERS,
    AlgorithmParams,
    ElementResult,
    algo_label,
    calculate_all_algorithms,
    num,
)

# ---------------------------------------------------------------------------
# Реквизиты документа (заполняются под конкретную организацию)
# ---------------------------------------------------------------------------

REPORT_META = {
    'organization': 'Гроднооблгаз',
    'department': 'Производственно-технический отдел',
    'city': 'г. Гродно',
    'doc_code': 'ГРО',
    'report_kind': 'ОТЧЁТ',
    'report_title': 'о результатах расчёта остаточного ресурса\n'
                    'оборудования газорегуляторного пункта',
}

# Цветовая схема
INK = (0, 0, 0)
GREY = (89, 89, 89)
RED = (176, 0, 0)

FILL_HEADER = 'D9E2F3'      # шапка таблицы
FILL_TOTAL = 'E8EDF5'        # чередование строк

BODY_FONT = 'Times New Roman'
BODY_SIZE = 14
TABLE_SIZE = 9


# ---------------------------------------------------------------------------
# Низкоуровневые помощники OOXML
# ---------------------------------------------------------------------------

def _shade(cell, fill):
    """Заливка ячейки таблицы."""
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement('w:shd')
    shading.set(qn('w:val'), 'clear')
    shading.set(qn('w:color'), 'auto')
    shading.set(qn('w:fill'), fill)
    properties.append(shading)


def _field(paragraph, instruction, size=BODY_SIZE, bold=False):
    """Вставляет поле Word (PAGE, NUMPAGES)."""
    run = paragraph.add_run()
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    begin = OxmlElement('w:fldChar')
    begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText')
    instr.set(qn('xml:space'), 'preserve')
    instr.text = instruction
    separate = OxmlElement('w:fldChar')
    separate.set(qn('w:fldCharType'), 'separate')
    placeholder = OxmlElement('w:t')
    placeholder.text = '1'
    end = OxmlElement('w:fldChar')
    end.set(qn('w:fldCharType'), 'end')
    for node in (begin, instr, separate, placeholder, end):
        run._r.append(node)
    return run


def _page_border(section, size='12', space='24'):
    """Рамка вокруг страницы — оформление титульного листа."""
    borders = OxmlElement('w:pgBorders')
    borders.set(qn('w:offsetFrom'), 'page')
    for edge in ('top', 'left', 'bottom', 'right'):
        element = OxmlElement(f'w:{edge}')
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), space)
        element.set(qn('w:color'), '000000')
        borders.append(element)
    section._sectPr.append(borders)


def _horizontal_rule(paragraph, size='8'):
    """Горизонтальная линия-разделитель (нижняя граница абзаца)."""
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size)
    bottom.set(qn('w:space'), '1')
    bottom.set(qn('w:color'), '808080')
    borders.append(bottom)
    properties.append(borders)


# ---------------------------------------------------------------------------
# Разметка страницы и стили
# ---------------------------------------------------------------------------

def _setup_section(section):
    """A4 с полями по ГОСТ: слева 30 мм, справа 15, сверху 20, снизу 20."""
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width, section.page_height = Cm(21.0), Cm(29.7)
    section.left_margin, section.right_margin = Cm(3.0), Cm(1.5)
    section.top_margin, section.bottom_margin = Cm(2.0), Cm(2.0)
    section.header_distance, section.footer_distance = Cm(1.0), Cm(1.0)


def _setup_styles(doc):
    """Единые стили документа."""
    normal = doc.styles['Normal']
    normal.font.name = BODY_FONT
    normal.font.size = Pt(BODY_SIZE)
    normal.font.color.rgb = RGBColor(*INK)
    paragraph_format = normal.paragraph_format
    paragraph_format.line_spacing = 1.5
    paragraph_format.space_after = Pt(0)
    paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

    for level, size, align, before, after in (
            (1, BODY_SIZE, WD_ALIGN_PARAGRAPH.CENTER, 18, 12),
            (2, 13, WD_ALIGN_PARAGRAPH.LEFT, 14, 8)):
        style = doc.styles[f'Heading {level}']
        style.font.name = BODY_FONT
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor(*INK)
        style.font.italic = False
        style.paragraph_format.alignment = align
        style.paragraph_format.space_before = Pt(before)
        style.paragraph_format.space_after = Pt(after)
        style.paragraph_format.keep_with_next = True
        style.paragraph_format.line_spacing = 1.15

    caption = doc.styles['Caption']
    caption.font.name = BODY_FONT
    caption.font.size = Pt(12)
    caption.font.bold = False
    caption.font.italic = True
    caption.font.color.rgb = RGBColor(*INK)
    caption.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    caption.paragraph_format.space_before = Pt(8)
    caption.paragraph_format.space_after = Pt(4)
    caption.paragraph_format.keep_with_next = True


# ---------------------------------------------------------------------------
# Текстовые блоки
# ---------------------------------------------------------------------------

def _body(doc, text, bold=False, color=INK, size=BODY_SIZE, align=None,
          indent=True, space_after=0, italic=False):
    paragraph = doc.add_paragraph()
    run = paragraph.add_run(text)
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    if align is not None:
        paragraph.alignment = align
    if indent and align is None:
        paragraph.paragraph_format.first_line_indent = Cm(1.25)
    paragraph.paragraph_format.space_after = Pt(space_after)
    return paragraph


def _note(doc, text, color=GREY, space_after=6):
    """Примечание меньшим кеглем."""
    return _body(doc, text, color=color, size=11, indent=False, italic=True,
                 space_after=space_after)


def _section_title(doc, text):
    return doc.add_heading(text, level=1)


# ---------------------------------------------------------------------------
# Таблицы
# ---------------------------------------------------------------------------

def _cell_text(cell, text, bold=False, size=TABLE_SIZE, color=INK,
               align=WD_ALIGN_PARAGRAPH.CENTER, fill=None, italic=False):
    cell.text = ''
    paragraph = cell.paragraphs[0]
    paragraph.alignment = align
    paragraph.paragraph_format.space_before = Pt(2)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    run = paragraph.add_run(str(text) if text is not None else '—')
    run.font.name = BODY_FONT
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = RGBColor(*color)
    if fill:
        _shade(cell, fill)


def _fix_widths(table, widths):
    """Фиксированная ширина колонок (иначе Word пересчитывает по содержимому)."""
    table.autofit = False
    layout = OxmlElement('w:tblLayout')
    layout.set(qn('w:type'), 'fixed')
    table._tbl.tblPr.append(layout)
    for row in table.rows:
        for index, width in enumerate(widths):
            if index < len(row.cells):
                row.cells[index].width = Cm(width)


def _repeat_header(row):
    """Повторять строку шапки на каждой странице."""
    properties = row._tr.get_or_add_trPr()
    header = OxmlElement('w:tblHeader')
    header.set(qn('w:val'), 'true')
    properties.append(header)


def _keep_together(row):
    """Запретить разрыв строки между страницами."""
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement('w:cantSplit')
    properties.append(cant_split)


def _table_caption(doc, number, title):
    paragraph = doc.add_paragraph(style='Caption')
    run = paragraph.add_run(f'Таблица {number} — {title}')
    run.font.name = BODY_FONT
    run.font.size = Pt(12)
    run.font.italic = True
    run.font.color.rgb = RGBColor(*INK)
    return paragraph


def _build_table(doc, headers, rows, widths, aligns=None, size=TABLE_SIZE,
                 header_size=None):
    """Таблица с повторяющейся шапкой и чередованием строк."""
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    table.style = 'Table Grid'
    table.alignment = WD_TABLE_ALIGNMENT.CENTER

    for index, header in enumerate(headers):
        _cell_text(table.rows[0].cells[index], header, bold=True,
                   size=header_size or size, fill=FILL_HEADER)
    _repeat_header(table.rows[0])
    _keep_together(table.rows[0])

    for row_index, row in enumerate(rows, start=1):
        fill = FILL_TOTAL if row_index % 2 == 0 else None
        for column, value in enumerate(row):
            align = (aligns[column] if aligns and column < len(aligns)
                     else WD_ALIGN_PARAGRAPH.CENTER)
            _cell_text(table.rows[row_index].cells[column], value, size=size,
                       align=align, fill=fill)
        _keep_together(table.rows[row_index])

    _fix_widths(table, widths)
    return table


# ---------------------------------------------------------------------------
# Титульный лист и подписи
# ---------------------------------------------------------------------------

def _title_page(doc, grp_name, grp_id, generated_at):
    """Титульный лист с рамкой и реквизитами."""
    section = doc.sections[0]
    _setup_section(section)
    _page_border(section)

    _body(doc, REPORT_META['organization'], bold=True, size=14,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)
    _body(doc, REPORT_META['department'], size=12,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, color=GREY)
    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(6)
    _horizontal_rule(rule)

    _body(doc, REPORT_META['report_kind'], bold=True, size=26,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False, space_after=10)
    for line in REPORT_META['report_title'].split('\n'):
        _body(doc, line, bold=True, size=16, align=WD_ALIGN_PARAGRAPH.CENTER,
              indent=False, space_after=2)

    for _ in range(4):
        doc.add_paragraph()

    details = [
        ('Объект оценки', f'ГРП «{grp_name}»'),
        ('Идентификатор объекта в системе', f'№ {grp_id}'),
        ('Шифр документа', REPORT_META['doc_code']),
        ('Редакция', '1.0'),
        ('Дата формирования', generated_at.strftime('%d.%m.%Y')),
        ('Режим доступа', 'Для служебного пользования'),
    ]
    table = doc.add_table(rows=len(details), cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for index, (label, value) in enumerate(details):
        _cell_text(table.rows[index].cells[0], label, align=WD_ALIGN_PARAGRAPH.RIGHT,
                   size=12, color=GREY)
        _cell_text(table.rows[index].cells[1], value, align=WD_ALIGN_PARAGRAPH.LEFT,
                   size=12, bold=True)
    _fix_widths(table, [8.0, 7.5])

    for _ in range(5):
        doc.add_paragraph()

    rule = doc.add_paragraph()
    rule.paragraph_format.space_before = Pt(6)
    rule.paragraph_format.space_after = Pt(6)
    _horizontal_rule(rule)

    _body(doc, REPORT_META['city'], size=13, align=WD_ALIGN_PARAGRAPH.CENTER,
          indent=False, space_after=2)
    _body(doc, str(generated_at.year), size=13, bold=True,
          align=WD_ALIGN_PARAGRAPH.CENTER, indent=False)


def _page_footer(section):
    """Колонтитул с номером страницы по центру (без наследования)."""
    section.footer.is_linked_to_previous = False
    footer = section.footer
    for paragraph in footer.paragraphs[1:]:
        paragraph._p.getparent().remove(paragraph._p)
    paragraph = footer.paragraphs[0]
    for run in list(paragraph.runs):
        run._r.getparent().remove(run._r)
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run('— ')
    run.font.name = BODY_FONT
    run.font.size = Pt(11)
    _field(paragraph, 'PAGE', size=11)
    run = paragraph.add_run(' —')
    run.font.name = BODY_FONT
    run.font.size = Pt(11)


def _signatures(doc, generated_at):
    """Подписи исполнителей."""
    _body(doc, '', indent=False)
    table = doc.add_table(rows=4, cols=3)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    rows = [
        ('Исполнитель', '', '_____________________'),
        ('Проверил', '', '_____________________'),
        ('Утвердил', '', '_____________________'),
        ('Дата выдачи', generated_at.strftime('%d.%m.%Y'), ''),
    ]
    for index, (label, value, signature) in enumerate(rows):
        _cell_text(table.rows[index].cells[0], label, align=WD_ALIGN_PARAGRAPH.LEFT,
                   size=12, bold=True)
        _cell_text(table.rows[index].cells[1], value, align=WD_ALIGN_PARAGRAPH.LEFT,
                   size=12, color=GREY)
        _cell_text(table.rows[index].cells[2], signature,
                   align=WD_ALIGN_PARAGRAPH.LEFT, size=12)
    _fix_widths(table, [4.5, 6.0, 5.0])


# ---------------------------------------------------------------------------
# Подготовка данных расчёта
# ---------------------------------------------------------------------------

def _match_journal(name, journal):
    """Записи журнала замен для типа оборудования."""
    low = str(name or '').strip().lower()
    for key, record in journal.items():
        if key and (key == low or key in low or low in key):
            return record
    return None


def _load_journal(db, grp_id):
    """Журнал замен, сгруппированный по типу оборудования.

    Отказы, повреждения и качество ремонта из журнала — вход расчёта
    (k_повр и K_рем), а не только украшение отчёта, поэтому читается всегда.
    """
    journal = {}
    try:
        rows = db.get_replacements_by_grp(grp_id) or []
    except Exception:  # noqa: BLE001 — журнал не критичен для отчёта
        return journal
    for row in rows:
        row = list(row) + [None] * 9
        haystack = ' '.join(str(v or '') for v in
                            (row[6], row[7], row[4], row[2])).lower()
        key = str(row[3] or '').strip().lower()
        record = journal.setdefault(key, {'fail': 0, 'damage': 0, 'poor': 0, 'total': 0})
        record['total'] += 1
        if any(marker in haystack for marker in FAILURE_MARKERS):
            record['fail'] += 1
        if any(marker in haystack for marker in DAMAGE_MARKERS):
            record['damage'] += 1
        if any(marker in haystack for marker in POOR_REPAIR_MARKERS):
            record['poor'] += 1
    return journal


def _calc_inputs(db, grp_id, equipment_rows, journal):
    """Входные данные для расчёта остаточного ресурса по алгоритмам."""
    payload = []
    for row in equipment_rows:
        row = list(row) + [None] * 4
        equip = {'name': row[1] or '', 'equipment_id': row[0],
                 'install_date': row[2], 'removal_date': row[3]}

        details = []
        try:
            parts = db.get_equipment_parts_full(row[0]) or []
        except Exception:  # noqa: BLE001
            parts = []
        for part in parts:
            part = list(part) + [None] * 8
            norm_years, removal, replaceable = part[3], part[5], part[7]
            if removal or not norm_years or float(norm_years) <= 0:
                continue
            details.append({'name': part[2] or 'Запчасть',
                            'norm_years': float(norm_years),
                            'install_date': part[4] or row[2],
                            'is_replaceable': bool(replaceable)})
        replaceable = [d for d in details if d['is_replaceable']]
        equip['details'] = replaceable or details

        record = _match_journal(equip['name'], journal)
        if record:
            equip['failures'] = record['fail']
            equip['damages'] = record['damage']
            if record['poor']:
                equip['repair'] = {'quality': 'poor'}
            elif record['total']:
                equip['repair'] = {'quality': 'new'}
        payload.append(equip)
    return payload


# ---------------------------------------------------------------------------
# Форма 6.2
# ---------------------------------------------------------------------------

# Методики, по которым печатается форма. В окне расчёта они называются
# «Алгоритм 3» и «Алгоритм 4» (algo_label), внутри расчёта — индексы 2 и 3.
FORM62_ALGORITHMS = (2, 3)

# Категории критических элементов в порядке формы 6.2:
# (ключ расчёта, слово в подписи, обозначение в форме).
FORM62_CATEGORIES = (
    ('regulator', 'регулятор', 'рег'),
    ('pzk', 'ПЗК', 'пзк'),
    ('psk', 'ПСК', 'пск'),
    ('filter', 'фильтр', 'фильтр'),
    ('valve', 'арматура', 'арматура'),
)

FORM62_KEYS = tuple(key for key, _word, _abbr in FORM62_CATEGORIES)

FORM62_HEADERS = ('Параметр', 'Обозначение', 'Значение', 'Ед. изм.', 'Примечание')

EMPTY = '—'


def _usable(result) -> bool:
    """Рассчитана ли методика: у неудачного расчёта заполнено поле error."""
    return result is not None and not result.error


def _trim_num(value, digits: int = 3) -> str:
    """Число без хвостовых нулей: «2,000» → «2», «10,000» → «10».

    Так печатается базовый ресурс: он приходит из норматива целым, и «2,000»
    в форме читалось бы как измеренное значение. Нераспознанное число («—»)
    возвращается как есть.
    """
    text = num(value, digits)
    if text == EMPTY or ',' not in text:
        return text
    return text.rstrip('0').rstrip(',') or text


def _z_text(value) -> str:
    """Ресурс в годах: три знака, хвостовые нули срезаны, но не все.

    «1,388» и «9,198» печатаются как есть, а «10,000» — как «10,0»: один знак
    остаётся всегда, иначе целое терялось бы среди дробных значений колонки.
    """
    text = num(value, 3)
    if text == EMPTY or ',' not in text:
        return text
    trimmed = text.rstrip('0')
    return trimmed if not trimmed.endswith(',') else trimmed + '0'


def _category_elements(result) -> Dict[str, ElementResult]:
    """По одному элементу на категорию формы 6.2 — определяющему её ресурс.

    Это элемент с наименьшим ресурсом (при равенстве — тот, чьё имя идёт
    раньше: тот же порядок, что у слабого звена в расчёте). Минимум по этим
    элементам равен ресурсу ГРП: в расчёт методики входят только критические
    элементы, и каждый попадает ровно в одну категорию.
    """
    chosen: Dict[str, ElementResult] = {}
    if not _usable(result):
        return chosen
    for element in result.elements:
        key = element.critical_key
        if not element.used or key not in FORM62_KEYS:
            continue
        current = chosen.get(key)
        if current is None or (element.z_element, element.name) < (
                current.z_element, current.name):
            chosen[key] = element
    return chosen


def _governing(chosen: Dict[str, ElementResult]) -> Optional[ElementResult]:
    """Элемент, задающий срок ГРП: минимум ресурса среди категорий."""
    if not chosen:
        return None
    return min(chosen.values(), key=lambda e: (e.z_element, e.name))


def _base_note(element: Optional[ElementResult]) -> str:
    """Откуда взят базовый ресурс элемента.

    ЭО — экспертная оценка. Приписка «min(Zбаза d)» ставится, когда база взята
    минимумом по заменяемым деталям: каскад вариантов расчёта строит такую базу
    формулой, начинающейся с «min(», и по ней это и видно.
    """
    if element is None:
        return EMPTY
    return 'ЭО min(Zбаза d)' if 'min(' in (element.z_base_formula or '') else 'ЭО'


def _product(element: Optional[ElementResult]) -> str:
    """Примечание к итоговому ресурсу: база и поправки произведением.

    Как в самой формуле расчёта: Z = Zбаза × Kсост × Kусл × Kрем × kповр.
    Базовый ресурс печатается без хвостовых нулей, коэффициент технического
    состояния — тремя знаками: он считается по состоянию элемента, и «1» на
    его месте скрывало бы, что состояние учитывалось.
    """
    if element is None:
        return EMPTY
    factors = (_trim_num(element.z_base), num(element.k_state, 3),
               num(element.k_cond, 1), num(element.k_repair, 1),
               num(element.k_fail, 1))
    return ' × '.join(factors)


def _min_note(chosen: Dict[str, ElementResult]) -> str:
    """Примечание строки «Слабое звено»: из чего выбран минимум."""
    values = [chosen[key].z_element for key in FORM62_KEYS if key in chosen]
    if not values:
        return EMPTY
    return 'min(' + '; '.join(_z_text(value) for value in values) + ')'


def _weak_value(result, chosen: Dict[str, ElementResult]) -> str:
    """Слабое звено словом — категория элемента с наименьшим ресурсом."""
    if not chosen:
        return EMPTY
    key = getattr(result, 'weak_element_key', None)
    if key not in FORM62_KEYS:
        # Категорию определяем сами: у методики слабое звено может быть не
        # записано (например, результат собран вне расчёта).
        key = _governing(chosen).critical_key
    for category, word, _abbr in FORM62_CATEGORIES:
        if category == key:
            return word
    return EMPTY


def form62_rows(result, params: AlgorithmParams = None,
                today=None) -> List[List[str]]:
    """Строки формы 6.2 по результату одной методики.

    Двадцать три строки — как в форме паспорта: базовые ресурсы, коэффициенты
    технического состояния, общие поправки, итоговые ресурсы, слабое звено,
    остаточный ресурс ГРП и срок следующего диагностирования. Шапку добавляет
    таблица отчёта.

    Если методика не рассчитана, строки остаются, а значения пусты: форма не
    должна пропадать из документа — причину печатает отчёт отдельной строкой.
    """
    params = params or AlgorithmParams()
    today = parse_date(today) or date.today()
    chosen = _category_elements(result)
    governing = _governing(chosen)
    rows: List[List[str]] = []

    # 1-5. Базовый ресурс: у элемента с деталями он взят минимумом по деталям.
    for key, word, abbr in FORM62_CATEGORIES:
        element = chosen.get(key)
        rows.append([f'Базовый ресурс ({word})', f'Zбаза,{abbr}',
                     EMPTY if element is None else _trim_num(element.z_base),
                     'лет', _base_note(element)])

    # 6-10. Коэффициент технического состояния считается по каждому элементу.
    for key, word, abbr in FORM62_CATEGORIES:
        element = chosen.get(key)
        rows.append([f'Коэффициент тех. состояния ({word})', f'Kсост,{abbr}',
                     EMPTY if element is None else num(element.k_state, 3),
                     EMPTY, 'Из расчетов' if element is not None else EMPTY])

    # 11-13. Поправки, одинаковые для всех элементов ГРП: они задаются на
    # объект целиком, поэтому в форме стоят по одной строке, а не по категориям.
    for label, abbr, attribute in (
            ('Коэффициент условий эксплуатации', 'Kусл', 'k_cond'),
            ('Коэффициент качества ремонта', 'Kрем', 'k_repair'),
            ('Коэффициент повреждений и отказов', 'kповр', 'k_fail')):
        rows.append([label, abbr,
                     EMPTY if governing is None
                     else num(getattr(governing, attribute), 1), EMPTY,
                     'Для всех элементов' if governing is not None else EMPTY])

    # 14-18. Итоговый ресурс элемента и его расчёт.
    for key, word, abbr in FORM62_CATEGORIES:
        element = chosen.get(key)
        rows.append([f'Итоговый ресурс ({word})', f'Z{abbr}',
                     EMPTY if element is None else _z_text(element.z_element),
                     'лет', _product(element)])

    # 19. Слабое звено — минимум по представителям категорий.
    rows.append(['Слабое звено', 'Zслабое', _weak_value(result, chosen),
                 EMPTY, _min_note(chosen)])

    # 20. Остаточный ресурс ГРП — ресурс слабого звена.
    rows.append(['Остаточный ресурс ГРП', 'ZГРП',
                 _z_text(result.result) if _usable(result) else EMPTY, 'лет',
                 'Ресурс слабого звена' if chosen else EMPTY])

    # 21-23. Срок следующего диагностирования и его дата.
    next_years = result.next_diagnosis if _usable(result) else None
    months = None if next_years is None else next_years * 12
    rows.append([
        'Срок следующего диагностирования', 'Tслед_диагн',
        EMPTY if next_years is None else num(next_years), 'лет',
        EMPTY if next_years is None else
        f'min({_z_text(result.result)} × {_trim_num(params.reserve, 2)}; '
        f'{_trim_num(params.max_diag_interval, 1)})'])
    rows.append([
        'Срок следующего диагностирования', 'Tслед_диагн',
        EMPTY if months is None else num(months, 2), 'мес.',
        EMPTY if months is None else f'{num(next_years)} × 12'])
    rows.append([
        'Рекомендуемая дата следующего диагностирования', EMPTY,
        month_year_text(expiry_date(next_years, today)), EMPTY,
        EMPTY if months is None else
        f'{today.strftime("%d.%m.%Y")} + {num(months, 2)} мес.'])
    return rows


# ---------------------------------------------------------------------------
# Отчёт
# ---------------------------------------------------------------------------

def generate_grp_docx(db, grp_id: int, grp_name: str, filename: str = None,
                      params: AlgorithmParams = None) -> str:
    """Сформировать Word-отчёт по конкретному ГРП.

    params — весовые коэффициенты алгоритмов. Если не переданы, берутся
    рекомендуемые значения методики (AlgorithmParams по умолчанию).
    Вызывающий код должен передавать те же параметры, что и в окно расчётов,
    иначе отчёт разойдётся с тем, что пользователь видел на экране.

    Возвращает путь к сохранённому файлу .docx.
    """
    generated_at = datetime.now()
    # Запасной кортеж короткий: сведения паспорта из него пусты, и таблица
    # «Сведения о ГРП» напечатает прочерки вместо падения.
    grp = db.get_grp_by_id(grp_id) or (grp_id, grp_name, 0, 0, 0)
    if params is None:
        params = AlgorithmParams()

    equipment_rows = db.get_equipment_by_grp(grp_id) or []
    journal = _load_journal(db, grp_id)
    results = calculate_all_algorithms(
        _calc_inputs(db, grp_id, equipment_rows, journal), params=params)

    doc = Document()
    _setup_styles(doc)

    # --- Титульный лист ---------------------------------------------------
    _title_page(doc, grp_name, grp_id, generated_at)

    # Разрыв раздела сам начинает новую страницу — лишний разрыв страницы
    # здесь создавал бы пустой лист.
    body_section = doc.add_section(WD_SECTION.NEW_PAGE)
    _setup_section(body_section)
    _page_footer(body_section)

    # --- 1. Сведения о ГРП -------------------------------------------------
    _section_title(doc, '1 СВЕДЕНИЯ О ГРП')
    _table_caption(doc, 1, 'Сведения о газорегуляторном пункте')
    _build_table(doc, ['Показатель', 'Значение'], report_rows(grp),
                 widths=[7.0, 9.5],
                 aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT],
                 size=11)
    _body(doc, '', indent=False, space_after=6)

    # --- 2. Форма 6.2 ------------------------------------------------------
    _section_title(doc, '2 РЕЗУЛЬТАТЫ РАСЧЁТА ОСТАТОЧНОГО РЕСУРСА')
    _body(doc, 'Результаты расчёта приведены по форме 6.2 для двух методик — '
               f'{algo_label(FORM62_ALGORITHMS[0])} и '
               f'{algo_label(FORM62_ALGORITHMS[1])}. Остаточный ресурс ГРП '
               'определён по слабому звену — минимальному ресурсу среди '
               'критических элементов.', space_after=6)

    for index, number in enumerate(FORM62_ALGORITHMS, start=1):
        result = results.get(number)
        doc.add_heading(f'2.{index} {algo_label(number)} — '
                        f'{ALGORITHM_TITLES.get(number, "")}', level=2)
        if not _usable(result):
            _note(doc, 'Расчёт по этой методике не выполнен: '
                       f'{result.error if result else "методика не рассчитана"}.',
                  color=RED)
        _table_caption(doc, 1 + index,
                       f'Результаты расчёта остаточного ресурса '
                       f'({algo_label(number)})')
        _build_table(doc, list(FORM62_HEADERS),
                     form62_rows(result, params, generated_at),
                     widths=[5.5, 2.6, 2.0, 1.4, 5.0],
                     aligns=[WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.CENTER,
                             WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.CENTER,
                             WD_ALIGN_PARAGRAPH.LEFT],
                     size=9)
        _body(doc, '', indent=False, space_after=6)

    _signatures(doc, generated_at)

    if not filename:
        filename = f'Отчёт_ГРП_{grp_id}_{generated_at:%Y%m%d}.docx'
    doc.save(filename)
    return filename
